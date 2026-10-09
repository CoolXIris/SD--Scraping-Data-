"""Scraper metadata anime Otakudesu dengan filter untuk aplikasi lokal."""

from __future__ import annotations

import argparse
import calendar
import csv
import re
import time
from dataclasses import dataclass, replace
from datetime import date, datetime
from pathlib import Path
from typing import Callable
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_URL = "https://otakudesu.blog"
CATEGORY_URLS = {
    "ongoing": f"{BASE_URL}/ongoing-anime/",
    "complete": f"{BASE_URL}/complete-anime/",
}
USER_AGENT = "LocalAnimeMetadata/1.0"
GENRE_LIST_URL = f"{BASE_URL}/genre-list/"
GENRE_OPTIONS_FALLBACK = [
    "Action", "Adventure", "Comedy", "Demons", "Drama", "Ecchi", "Fantasy",
    "Game", "Harem", "Historical", "Horror", "Josei", "Magic", "Martial Arts",
    "Mecha", "Military", "Music", "Mystery", "Parody", "Police", "Romance",
    "Samurai", "School", "Sci-Fi", "Seinen", "Shoujo", "Shoujo Ai", "Shounen",
    "Slice of Life", "Space", "Sports", "Super Power", "Supernatural", "Thriller",
    "Vampire",
]
MONTHS = {
    "jan": 1, "januari": 1, "feb": 2, "februari": 2,
    "mar": 3, "maret": 3, "apr": 4, "april": 4,
    "may": 5, "mei": 5, "jun": 6, "juni": 6,
    "jul": 7, "juli": 7, "aug": 8, "agu": 8, "agustus": 8,
    "sep": 9, "september": 9, "oct": 10, "okt": 10, "oktober": 10,
    "nov": 11, "november": 11, "dec": 12, "des": 12, "desember": 12,
}
WEEKDAY_MAP = {
    "senin": "Senin", "monday": "Senin",
    "selasa": "Selasa", "tuesday": "Selasa",
    "rabu": "Rabu", "wednesday": "Rabu",
    "kamis": "Kamis", "thursday": "Kamis",
    "jumat": "Jumat", "friday": "Jumat",
    "sabtu": "Sabtu", "saturday": "Sabtu",
    "minggu": "Minggu", "sunday": "Minggu",
}
AnimeProgress = Callable[[str, int], None]


@dataclass(frozen=True)
class AnimeRelease:
    title: str
    date: date
    weekday: str
    url: str
    status: str = "Ongoing"
    episode: str = ""
    genres: tuple[str, ...] = ()


def subtract_months(value: date, months: int) -> date:
    month_index = value.year * 12 + value.month - 1 - months
    year, zero_based_month = divmod(month_index, 12)
    month = zero_based_month + 1
    day = min(value.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def parse_listing_date(value: str, reference_date: date) -> date | None:
    match = re.search(
        r"(?P<day>\d{1,2})\s+(?P<month>[A-Za-z]+)(?:\s+(?P<year>\d{4}))?",
        value,
        re.IGNORECASE,
    )
    if not match:
        return None

    month = MONTHS.get(match.group("month").lower())
    if month is None:
        return None

    day = int(match.group("day"))
    explicit_year = match.group("year")
    years = [int(explicit_year)] if explicit_year else [
        reference_date.year - 1,
        reference_date.year,
        reference_date.year + 1,
    ]
    candidates = []
    for year in years:
        try:
            candidate = date(year, month, day)
        except ValueError:
            continue
        if candidate <= reference_date:
            candidates.append(candidate)
    return max(candidates, default=None)


def episode_label(item: Tag) -> str:
    episode = item.select_one(".epz")
    return episode.get_text(" ", strip=True) if episode else ""


def episode_number(item: Tag) -> int | None:
    text = episode_label(item)
    match = re.search(r"episode\s*(\d+)|(\d+)\s*episode", text, re.IGNORECASE)
    if not match:
        return None
    return int(match.group(1) or match.group(2))


def normalize_weekday(value: str) -> str:
    key = re.sub(r"[^a-z]", "", value.casefold())
    return WEEKDAY_MAP.get(key, "")


def parse_anime_item(
    item: Tag,
    reference_date: date,
    status: str = "ongoing",
) -> tuple[AnimeRelease | None, date | None]:
    date_element = item.select_one(".newnime")
    if date_element is None:
        return None, None
    posted_date = parse_listing_date(date_element.get_text(" ", strip=True), reference_date)
    if posted_date is None:
        return None, None

    title_element = item.select_one(".jdlflm")
    if title_element is not None:
        title = title_element.get_text(" ", strip=True)
    else:
        image = item.select_one(".thumb img[alt]")
        title = image.get("alt", "").strip() if image else ""
        title = re.sub(r"\s+Sub Indo$", "", title, flags=re.IGNORECASE)

    anime_link = item.select_one(".thumb a[href], .jdlflm a[href]")
    if not title or anime_link is None:
        return None, posted_date

    weekday_element = item.select_one(".epztipe")
    weekday_text = weekday_element.get_text(" ", strip=True) if weekday_element else ""
    weekday = normalize_weekday(weekday_text) if status == "ongoing" else ""
    status_label = "Ongoing" if status == "ongoing" else "Complete"
    return AnimeRelease(
        title=title,
        date=posted_date,
        weekday=weekday,
        url=urljoin(BASE_URL, str(anime_link["href"])),
        status=status_label,
        episode=episode_label(item),
    ), posted_date


def next_page_number(soup: BeautifulSoup, current_page: int, status: str = "ongoing") -> int | None:
    route = "ongoing-anime" if status == "ongoing" else "complete-anime"
    page_numbers = set()
    for link in soup.select(".pagination a[href]"):
        page_url = urljoin(BASE_URL, str(link["href"]))
        if urlparse(page_url).netloc != urlparse(BASE_URL).netloc:
            continue
        match = re.search(rf"/{route}/page/(\d+)/?", urlparse(page_url).path)
        if match and int(match.group(1)) > current_page:
            page_numbers.add(int(match.group(1)))
    return min(page_numbers, default=None)


def fetch_genre_options(timeout: int = 15) -> list[str]:
    response = requests.get(
        GENRE_LIST_URL,
        headers={"User-Agent": USER_AGENT},
        timeout=timeout,
    )
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    genres = sorted({
        link.get_text(" ", strip=True)
        for link in soup.select('a[href*="/genres/"]')
        if link.get_text(" ", strip=True)
    }, key=str.casefold)
    return genres or GENRE_OPTIONS_FALLBACK.copy()


def fetch_anime_genres(session: requests.Session, url: str, delay: float = 0) -> tuple[str, ...]:
    response = session.get(url, timeout=20)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    genre_links = soup.select('.infozingle a[rel="tag"], .infozingle a[href*="/genres/"]')
    genres = tuple(dict.fromkeys(
        link.get_text(" ", strip=True)
        for link in genre_links
        if link.get_text(" ", strip=True)
    ))
    if delay:
        time.sleep(delay)
    return genres


def _report(progress: AnimeProgress | None, message: str, percent: int) -> None:
    if progress:
        progress(message, max(0, min(99, percent)))


def _collect_status(
    session: requests.Session,
    status: str,
    cutoff: date,
    today: date,
    release_days: set[str],
    max_pages: int,
    delay: float,
    progress: AnimeProgress | None,
    status_index: int,
    status_count: int,
) -> list[AnimeRelease]:
    base_url = CATEGORY_URLS[status]
    page_number = 1
    visited_pages: set[int] = set()
    releases: list[AnimeRelease] = []

    while page_number not in visited_pages and (not max_pages or page_number <= max_pages):
        visited_pages.add(page_number)
        page_url = base_url if page_number == 1 else f"{base_url}page/{page_number}/"
        response = session.get(page_url, timeout=25)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        items = soup.select(".venz li")
        if not items:
            if page_number == 1:
                raise RuntimeError(f"Daftar anime {status} tidak ditemukan (.venz li).")
            break

        page_dates: list[date] = []
        for item in items:
            release, posted_date = parse_anime_item(item, today, status)
            if posted_date is not None:
                page_dates.append(posted_date)
            if release and cutoff <= release.date <= today:
                if status == "ongoing" and release_days and release.weekday not in release_days:
                    continue
                releases.append(release)

        base_percent = int((status_index / max(status_count, 1)) * 40)
        page_percent = min(page_number - 1, 20)
        _report(progress, f"Membaca daftar {status_label(status)} — halaman {page_number}", 8 + base_percent + page_percent)

        if page_dates and max(page_dates) < cutoff:
            break

        next_page = next_page_number(soup, page_number, status)
        if next_page is None:
            break
        page_number = next_page
        if delay:
            time.sleep(delay)

    return releases


def status_label(status: str) -> str:
    return "ongoing" if status == "ongoing" else "complete"


def scrape_anime(
    months: int = 2,
    status: str = "ongoing",
    genre: str = "",
    release_days: list[str] | set[str] | tuple[str, ...] = (),
    max_pages: int = 0,
    delay: float = 0.75,
    progress: AnimeProgress | None = None,
    reference_date: date | None = None,
) -> list[AnimeRelease]:
    if months < 1:
        raise ValueError("Rentang waktu anime harus minimal 1 bulan.")
    if status not in {"ongoing", "complete", "both"}:
        raise ValueError("Status anime harus ongoing, complete, atau both.")
    if max_pages < 0 or delay < 0:
        raise ValueError("Batas halaman dan jeda tidak boleh negatif.")

    today = reference_date or date.today()
    cutoff = subtract_months(today, months)
    selected_days = {normalize_weekday(day) for day in release_days if normalize_weekday(day)}
    statuses = ["ongoing", "complete"] if status == "both" else [status]
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})

    candidates: list[AnimeRelease] = []
    try:
        for index, category in enumerate(statuses):
            candidates.extend(_collect_status(
                session=session,
                status=category,
                cutoff=cutoff,
                today=today,
                release_days=selected_days,
                max_pages=max_pages,
                delay=delay,
                progress=progress,
                status_index=index,
                status_count=len(statuses),
            ))

        # If the same series appears in both categories, keep the first listing
        # (ongoing is collected first) so one anime does not appear twice.
        unique: dict[str, AnimeRelease] = {}
        for release in candidates:
            unique.setdefault(release.url.rstrip("/").casefold(), release)
        candidates = list(unique.values())

        selected_genre = genre.strip().casefold()
        if selected_genre:
            matched: list[AnimeRelease] = []
            total = len(candidates)
            for index, release in enumerate(candidates, start=1):
                if index > 1 and delay:
                    time.sleep(delay)
                try:
                    genres = fetch_anime_genres(session, release.url)
                except requests.RequestException as error:
                    raise RuntimeError(f"Gagal membaca genre untuk {release.title}: {error}") from error
                if any(value.casefold() == selected_genre for value in genres):
                    matched.append(replace(release, genres=genres))
                if index % 5 == 0 or index == total:
                    percent = 55 + int(index / max(total, 1) * 40)
                    _report(progress, f"Memeriksa genre {index}/{total}", percent)
            candidates = matched

        _report(progress, "Menyiapkan hasil anime", 98)
        return sorted(candidates, key=lambda item: (item.date, item.title.casefold()), reverse=True)
    finally:
        session.close()


def scrape_episode_ones(
    months: int,
    max_pages: int,
    delay: float,
    reference_date: date | None = None,
) -> list[AnimeRelease]:
    """Kompatibilitas untuk scraper lama: anime ongoing yang memulai Episode 1."""
    releases = scrape_anime(
        months=months,
        status="ongoing",
        max_pages=max_pages,
        delay=delay,
        reference_date=reference_date,
    )
    return [release for release in releases if episode_number_from_label(release.episode) == 1]


def episode_number_from_label(value: str) -> int | None:
    match = re.search(r"episode\s*(\d+)|(\d+)\s*episode", value, re.IGNORECASE)
    return int(match.group(1) or match.group(2)) if match else None


def save_csv(releases: list[AnimeRelease], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as output_file:
        writer = csv.DictWriter(
            output_file,
            fieldnames=["judul", "status", "episode", "genre", "tanggal_post", "hari_rilis", "halaman_anime"],
        )
        writer.writeheader()
        for release in releases:
            writer.writerow({
                "judul": release.title,
                "status": release.status,
                "episode": release.episode,
                "genre": ", ".join(release.genres),
                "tanggal_post": release.date.isoformat(),
                "hari_rilis": release.weekday,
                "halaman_anime": release.url,
            })


def save_anime_xlsx(releases: list[AnimeRelease], output_path: Path) -> None:
    columns = ["Judul", "Status", "Episode", "Genre", "Tanggal Posting", "Hari Rilis", "Halaman Anime"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Anime"
    sheet.sheet_view.showGridLines = False
    sheet.append(columns)
    for release in releases:
        sheet.append([
            release.title,
            release.status,
            release.episode,
            ", ".join(release.genres),
            datetime.combine(release.date, datetime.min.time()),
            release.weekday,
            release.url,
        ])

    header_fill = PatternFill(fill_type="solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 28
    sheet.auto_filter.ref = f"A1:G{max(len(releases) + 1, 1)}"
    for column, width in zip("ABCDEFG", [48, 14, 18, 42, 20, 16, 60]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[4].number_format = "yyyy-mm-dd"
        if row[6].value:
            row[6].hyperlink = row[6].value
            row[6].style = "Hyperlink"
    if releases:
        table = Table(displayName="TabelAnime", ref=f"A1:G{len(releases) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Ambil metadata anime Otakudesu dengan filter status, waktu, genre, dan hari rilis.")
    parser.add_argument("--months", type=int, choices=(1, 6, 12), default=1, help="Rentang tanggal yang dihitung ke belakang.")
    parser.add_argument("--status", choices=("ongoing", "complete", "both"), default="ongoing")
    parser.add_argument("--genre", default="", help="Genre yang harus dimiliki anime.")
    parser.add_argument("--release-day", action="append", default=[], help="Hari rilis ongoing; bisa dipakai beberapa kali.")
    parser.add_argument("--max-pages", type=int, default=0, help="Batas halaman per status; 0 berarti ikuti pagination sampai batas tanggal.")
    parser.add_argument("--delay", type=float, default=0.75)
    parser.add_argument("--output", type=Path, default=Path("anime_terbaru.xlsx"))
    args = parser.parse_args()
    if args.max_pages < 0 or args.delay < 0:
        parser.error("max-pages dan delay tidak boleh negatif")

    try:
        releases = scrape_anime(
            months=args.months,
            status=args.status,
            genre=args.genre,
            release_days=args.release_day,
            max_pages=args.max_pages,
            delay=args.delay,
            progress=lambda message, _: print(message),
        )
        save_anime_xlsx(releases, args.output)
    except (requests.RequestException, RuntimeError, OSError, ValueError) as error:
        parser.exit(1, f"Gagal mengambil atau menyimpan data anime: {error}\n")

    print(f"Anime ditemukan: {len(releases)}")
    print(f"File XLSX tersimpan di: {args.output.resolve()}")


if __name__ == "__main__":
    main()

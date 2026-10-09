"""Buat template CSV dan ekspor data lowongan ke Excel (.xlsx).

Isi CSV secara manual atau gunakan file ekspor yang memang diizinkan sumbernya,
kemudian jalankan perintah ``export`` untuk membuat workbook Excel.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import TextIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMNS = [
    "Jabatan",
    "Perusahaan",
    "Deskripsi Pekerjaan",
    "Kualifikasi",
    "Gaji",
    "Lokasi",
    "Tanggal Posting",
    "URL Lowongan",
]

ALIASES = {
    "Jabatan": {
        "jabatan", "jabatanlowongan", "jabatanlowonganyangdibuka", "posisi",
        "judullowongan", "title", "jobtitle", "role",
    },
    "Perusahaan": {"perusahaan", "namaperusahaan", "company", "companyname", "employer"},
    "Deskripsi Pekerjaan": {
        "deskripsi", "deskripsipekerjaan", "deskripsipekerjaannya",
        "jobdescription", "description",
    },
    "Kualifikasi": {"kualifikasi", "persyaratan", "requirements", "qualification", "qualifications"},
    "Gaji": {"gaji", "salary", "pay", "remuneration"},
    "Lokasi": {"lokasi", "location", "joblocation"},
    "Tanggal Posting": {
        "tanggal", "tanggalposting", "tanggaldibuat", "tanggallowongandibuat",
        "dateposted", "postingdate", "createddate",
    },
    "URL Lowongan": {"url", "link", "urllowongan", "linklowongan", "joburl"},
}


def normalize_header(value: str) -> str:
    """Samakan variasi kapitalisasi, spasi, dan tanda baca pada nama kolom."""
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def open_csv(path: Path) -> tuple[TextIO, csv.DictReader]:
    """Buka CSV UTF-8 (termasuk UTF-8 BOM) dan deteksi pemisah umum."""
    source = path.open("r", encoding="utf-8-sig", newline="")
    sample = source.read(4096)
    source.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return source, csv.DictReader(source, dialect=dialect)


def column_mapping(fieldnames: list[str] | None) -> dict[str, str]:
    if not fieldnames:
        raise ValueError("CSV tidak memiliki baris header.")

    normalized_aliases = {
        column: {normalize_header(alias) for alias in aliases}
        for column, aliases in ALIASES.items()
    }
    mapping: dict[str, str] = {}
    for source_name in fieldnames:
        normalized = normalize_header(source_name or "")
        for output_name, aliases in normalized_aliases.items():
            if output_name not in mapping and normalized in aliases:
                mapping[output_name] = source_name
                break

    if "Jabatan" not in mapping:
        raise ValueError(
            "Kolom jabatan tidak ditemukan. Gunakan header 'Jabatan', 'Posisi', "
            "'Judul Lowongan', atau 'Job Title'."
        )
    return mapping


def create_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=COLUMNS)
        writer.writeheader()


def export_xlsx(input_path: Path, output_path: Path) -> int:
    source, reader = open_csv(input_path)
    try:
        mapping = column_mapping(reader.fieldnames)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Lowongan Sumsel"
        sheet.append(COLUMNS)

        row_count = 0
        for source_row in reader:
            if not any((value or "").strip() for value in source_row.values() if isinstance(value, str)):
                continue
            sheet.append([
                (source_row.get(mapping[column], "") or "").strip()
                if column in mapping
                else ""
                for column in COLUMNS
            ])
            row_count += 1
    finally:
        source.close()

    header_fill = PatternFill(fill_type="solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30
    sheet.auto_filter.ref = f"A1:H{max(row_count + 1, 1)}"
    widths = [28, 28, 55, 45, 25, 30, 20, 48]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if row_count:
        table = Table(displayName="TabelLowonganSumsel", ref=f"A1:H{row_count + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return row_count


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Buat template CSV atau ekspor data lowongan Sumatera Selatan ke XLSX."
    )
    commands = parser.add_subparsers(dest="command", required=True)

    template_parser = commands.add_parser("template", help="Buat CSV kosong dengan kolom yang diperlukan.")
    template_parser.add_argument(
        "--output",
        type=Path,
        default=Path("lowongan_sumsel_template.csv"),
        help="Lokasi file CSV template.",
    )

    export_parser = commands.add_parser("export", help="Konversi CSV menjadi workbook XLSX.")
    export_parser.add_argument("input", type=Path, help="File CSV sumber.")
    export_parser.add_argument(
        "--output",
        type=Path,
        default=Path("lowongan_sumatera_selatan.xlsx"),
        help="Lokasi file XLSX keluaran.",
    )

    args = parser.parse_args()
    try:
        if args.command == "template":
            create_template(args.output)
            print(f"Template CSV dibuat: {args.output.resolve()}")
        else:
            if not args.input.is_file():
                parser.error(f"File CSV tidak ditemukan: {args.input}")
            count = export_xlsx(args.input, args.output)
            print(f"{count} lowongan diekspor ke: {args.output.resolve()}")
    except (OSError, ValueError, csv.Error) as error:
        parser.exit(1, f"Gagal memproses file: {error}\n")


if __name__ == "__main__":
    main()

"""Pengambil tabel statistik CSA dari WebAPI resmi BPS."""

from __future__ import annotations

import argparse
import re
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Callable, Iterator
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


API_ROOT = "https://webapi.bps.go.id/v1/api"
API_LIST = f"{API_ROOT}/list"
API_VIEW = f"{API_ROOT}/view"
USER_AGENT = "ScraperDataLocal/1.0"
ProgressCallback = Callable[[str, int], None]

# IDs kategori CSA yang ditampilkan pada halaman Statistik menurut Subjek BPS.
SUBJECT_CATEGORIES = [
    {"id": "514", "name": "Statistik Demografi & Sosial"},
    {"id": "515", "name": "Statistik Ekonomi"},
    {"id": "516", "name": "Statistik Lingkungan Hidup dan Multi-domain"},
]
TABLE_SOURCES = {"1": "Tabel Statistik", "2": "Tabel Dinamis", "3": "SIMDASI"}
OUTPUT_COLUMNS = ["Tahun", "Kategori Subjek", "Jenis Data", "Nama Tabel", "Lembar", "Variabel", "Nilai"]


class BPSAPIError(RuntimeError):
    """Kesalahan WebAPI BPS tanpa membocorkan API key pada pesan URL."""


def _check_key(api_key: str) -> str:
    key = str(api_key or "").strip()
    if not key:
        raise ValueError("Masukkan API key WebAPI BPS terlebih dahulu.")
    return key


def _request_json(session: requests.Session, url: str, params: dict, api_key: str) -> dict:
    query = {**params, "key": _check_key(api_key)}
    try:
        response = session.get(url, params=query, timeout=30)
        response.raise_for_status()
    except requests.RequestException as error:
        raise BPSAPIError(f"Permintaan ke WebAPI BPS gagal ({type(error).__name__}). Periksa koneksi dan API key.") from None

    try:
        payload = response.json()
    except ValueError:
        raise BPSAPIError("Respons WebAPI BPS bukan JSON. Periksa API key dan coba lagi.") from None
    if not isinstance(payload, dict):
        raise BPSAPIError("Format respons WebAPI BPS tidak sesuai.")
    if str(payload.get("status", "OK")).casefold() not in {"ok", "200", "success"}:
        message = str(payload.get("message") or payload.get("error") or "WebAPI BPS menolak permintaan.")
        message = message.replace(query["key"], "[API key disembunyikan]")
        raise BPSAPIError(message[:250])
    return payload


def _list_from_payload(payload: dict) -> tuple[list[dict], dict]:
    data = payload.get("data")
    if isinstance(data, list):
        meta = data[0] if data and isinstance(data[0], dict) else {}
        for candidate in data[1:]:
            if isinstance(candidate, list):
                return [item for item in candidate if isinstance(item, dict)], meta
        return [], meta
    if isinstance(data, dict):
        inner = data.get("data")
        if isinstance(inner, list):
            return [item for item in inner if isinstance(item, dict)], data
    return [], {}


def _api_list(session: requests.Session, api_key: str, model: str, **params) -> tuple[list[dict], dict]:
    payload = _request_json(
        session,
        API_LIST,
        {"model": model, "domain": "0000", "lang": "ind", **params},
        api_key,
    )
    return _list_from_payload(payload)


def _paginated_list(session: requests.Session, api_key: str, model: str, **params) -> list[dict]:
    results: list[dict] = []
    seen: set[str] = set()
    page = 1
    while page <= 1000:
        records, meta = _api_list(session, api_key, model, page=page, **params)
        for record in records:
            identity = str(record.get("id") or record.get("sub_id") or record.get("subcat_id") or record.get("title") or "")
            if identity and identity not in seen:
                seen.add(identity)
                results.append(record)
        try:
            pages = int(meta.get("pages", page))
        except (TypeError, ValueError):
            pages = page
        if page >= pages or not records:
            break
        page += 1
    return results


def fetch_subject_types(api_key: str, category_id: str) -> list[dict[str, str]]:
    """Ambil jenis/subjek statistik di dalam satu kategori CSA."""
    if str(category_id) not in {item["id"] for item in SUBJECT_CATEGORIES}:
        raise ValueError("Kategori subjek BPS tidak dikenal.")
    with requests.Session() as session:
        session.headers.update({"Accept": "application/json", "User-Agent": USER_AGENT})
        subjects = _paginated_list(session, api_key, "subjectcsa", subcat=str(category_id))
    return sorted(
        [
            {"id": str(item.get("sub_id") or ""), "name": str(item.get("title") or "").strip()}
            for item in subjects
            if item.get("sub_id") and item.get("title")
        ],
        key=lambda item: item["name"].casefold(),
    )


def fetch_tables(api_key: str, subject_id: str) -> list[dict[str, str]]:
    """Ambil tabel statistik pada satu jenis/subjek CSA."""
    subject_id = str(subject_id or "").strip()
    if not subject_id.isdigit():
        raise ValueError("Pilih jenis data BPS terlebih dahulu.")
    with requests.Session() as session:
        session.headers.update({"Accept": "application/json", "User-Agent": USER_AGENT})
        tables = _paginated_list(session, api_key, "tablestatistic", subject=subject_id, perpage=100)
    return sorted(
        [
            {
                "id": str(item.get("id") or ""),
                "title": str(item.get("title") or "").strip(),
                "source": str(item.get("tablesource") or ""),
                "source_name": TABLE_SOURCES.get(str(item.get("tablesource") or ""), "Tabel Statistik"),
            }
            for item in tables
            if item.get("id") and item.get("title")
        ],
        key=lambda item: item["title"].casefold(),
    )


def _safe_excel_url(value: object) -> str:
    url = str(value or "").strip()
    if url.startswith("//"):
        url = f"https:{url}"
    elif url.startswith("/"):
        url = urljoin("https://web-api.bps.go.id/", url)
    elif url and not urlparse(url).scheme:
        url = urljoin("https://web-api.bps.go.id/", url)
    parsed = urlparse(url)
    host = (parsed.hostname or "").casefold()
    if parsed.scheme != "https" or not (host == "bps.go.id" or host.endswith(".bps.go.id")):
        return ""
    return url


def _year_number(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    match = re.fullmatch(r"\s*(\d{4})(?:\.0)?\s*", str(value))
    return int(match.group(1)) if match else None


def _xlsx_rows_for_years(content: bytes, start_year: int, end_year: int, fallback_year: int) -> list[list[object]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    output: list[list[object]] = []
    try:
        for sheet in workbook.worksheets:
            grid = [list(row) for row in sheet.iter_rows(values_only=True)]
            if not grid:
                continue

            # BPS workbooks usually place years down the first column or across
            # a header row. Convert either layout to one row per year/value.
            row_years = [
                (row_index, _year_number(row[0]) if row else None)
                for row_index, row in enumerate(grid)
            ]
            matching_rows = [(idx, year) for idx, year in row_years if year is not None and start_year <= year <= end_year]
            if matching_rows:
                first_year_index = min(idx for idx, _ in matching_rows)
                header = grid[max(0, first_year_index - 1)] if first_year_index else grid[0]
                for row_index, year in matching_rows:
                    row = grid[row_index]
                    fields = [
                        f"{header[col] if col < len(header) and header[col] not in (None, '') else f'Kolom {col + 1}'}: {value}"
                        for col, value in enumerate(row[1:], start=1)
                        if value not in (None, "")
                    ]
                    for field in fields:
                        label, _, value = field.partition(": ")
                        output.append([year, sheet.title, label, value])
                continue

            year_columns: dict[int, tuple[int, int]] = {}
            for row_index, row in enumerate(grid[:8]):
                for column_index, value in enumerate(row):
                    year = _year_number(value)
                    if year is not None and start_year <= year <= end_year:
                        year_columns[year] = (row_index, column_index)
            if year_columns:
                first_header = min(row_index for row_index, _ in year_columns.values())
                for year, (_, year_column) in sorted(year_columns.items()):
                    for row in grid[first_header + 1:]:
                        if year_column >= len(row) or row[year_column] in (None, ""):
                            continue
                        label = " | ".join(str(cell).strip() for cell in row[:year_column] if cell not in (None, ""))
                        output.append([year, sheet.title, label or "Data", row[year_column]])
                continue

            for row in grid:
                cells = [str(value).strip() for value in row if value not in (None, "")]
                if cells:
                    output.append([fallback_year, sheet.title, "Data", " | ".join(cells)])
    finally:
        workbook.close()
    return output


def _iter_leaves(value: object, prefix: str = "") -> Iterator[tuple[str, object]]:
    if isinstance(value, dict):
        for key, child in value.items():
            key_text = str(key)
            if key_text.casefold() in {
                "excel", "pdf", "image", "url", "id", "title", "year",
                "status", "data-availability", "updt_date", "oldest_period",
                "latest_period", "size", "tablesource", "subject", "subcat",
            }:
                continue
            next_prefix = f"{prefix}.{key_text}" if prefix else key_text
            yield from _iter_leaves(child, next_prefix)
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            yield from _iter_leaves(child, f"{prefix}[{index}]")
    else:
        yield prefix or "Data", value


def scrape_table_by_year(
    api_key: str,
    category_name: str,
    subject_name: str,
    table_id: str,
    table_title: str,
    table_source: str,
    start_year: int,
    end_year: int,
    delay: float = 0.25,
    progress: ProgressCallback | None = None,
) -> tuple[list[list[object]], int]:
    """Ambil detail tabel untuk setiap tahun dan ratakan hasil untuk XLSX."""
    _check_key(api_key)
    if not table_id:
        raise ValueError("Pilih nama tabel statistik terlebih dahulu.")
    current_year = datetime.now().year
    if not (1900 <= start_year <= end_year <= current_year):
        raise ValueError(f"Rentang tahun harus di antara 1900 dan {current_year}.")
    if end_year - start_year > 100:
        raise ValueError("Rentang tahun maksimal 101 tahun per proses.")
    if delay < 0:
        raise ValueError("Jeda permintaan tidak boleh negatif.")

    records: list[list[object]] = []
    unavailable_years = 0
    cached_excel: dict[str, list[list[object]]] = {}
    years = list(range(start_year, end_year + 1))
    session = requests.Session()
    session.headers.update({"Accept": "application/json", "User-Agent": USER_AGENT})
    try:
        for index, year in enumerate(years, start=1):
            payload = _request_json(
                session,
                API_VIEW,
                {
                    "model": "tablestatistic",
                    "domain": "0000",
                    "id": table_id,
                    "year": year,
                    "lang": "ind",
                },
                api_key,
            )
            if str(payload.get("data-availability", "available")).casefold() != "available":
                unavailable_years += 1
                if progress:
                    progress(f"Tidak ada data BPS untuk tahun {year} ({index}/{len(years)})", 10 + int(index / len(years) * 86))
                if delay and index < len(years):
                    time.sleep(delay)
                continue
            table_data = payload.get("data")
            if isinstance(table_data, list) and table_data:
                detail = next((item for item in table_data if isinstance(item, dict) and any(
                    key in item for key in ("excel", "datacontent", "table", "rows", "values")
                )), None)
                table_data = detail or next((item for item in table_data if isinstance(item, dict)), table_data)
            if not isinstance(table_data, (dict, list)):
                table_data = payload

            excel_url = ""
            if isinstance(table_data, dict):
                excel_url = _safe_excel_url(table_data.get("excel"))
            year_rows: list[list[object]] = []
            if excel_url:
                if excel_url not in cached_excel:
                    try:
                        file_response = session.get(excel_url, timeout=45)
                        file_response.raise_for_status()
                        cached_excel[excel_url] = _xlsx_rows_for_years(
                            file_response.content, start_year, end_year, year
                        )
                    except (requests.RequestException, ValueError, OSError) as error:
                        raise BPSAPIError(f"File tabel BPS gagal dibaca ({type(error).__name__}).") from None
                year_rows = [row for row in cached_excel[excel_url] if row[0] == year]

            if not year_rows:
                content = table_data
                if isinstance(content, dict):
                    for field in ("datacontent", "table", "rows", "values", "data"):
                        if field in content and content[field] not in (None, "", [], {}):
                            content = content[field]
                            break
                if isinstance(content, (dict, list)):
                    year_rows = [[year, "API", name, value] for name, value in _iter_leaves(content)]

            if year_rows:
                for row in year_rows:
                    row_year, sheet_name, variable, value = row[:4]
                    records.append([row_year, category_name, subject_name, table_title, sheet_name, variable, value])
            else:
                unavailable_years += 1

            if progress:
                progress(f"Mengambil data BPS tahun {year} ({index}/{len(years)})", 10 + int(index / len(years) * 86))
            if delay and index < len(years):
                time.sleep(delay)
    finally:
        session.close()

    return records, unavailable_years


def save_xlsx(rows: list[list[object]], output_path: Path, metadata: dict[str, object] | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data BPS"
    sheet.sheet_view.showGridLines = False
    sheet.append(OUTPUT_COLUMNS)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill(fill_type="solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{max(len(rows) + 1, 1)}"
    sheet.row_dimensions[1].height = 30
    for column, width in zip("ABCDEFG", [12, 38, 32, 64, 22, 58, 34]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        table = Table(displayName="TabelStatistikBPS", ref=f"A1:G{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)

    if metadata:
        info = workbook.create_sheet("Ringkasan")
        info.append(["Informasi", "Nilai"])
        for key, value in metadata.items():
            info.append([key, value])
        for cell in info[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        info.column_dimensions["A"].width = 28
        info.column_dimensions["B"].width = 80
        info.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Ambil tabel statistik berdasarkan subjek CSA BPS.")
    parser.add_argument("--api-key", required=True, help="API key dari portal WebAPI BPS.")
    parser.add_argument("--category-id", required=True, choices=[item["id"] for item in SUBJECT_CATEGORIES])
    parser.add_argument("--subject-id", required=True)
    parser.add_argument("--table-id", required=True)
    parser.add_argument("--title", required=True)
    parser.add_argument("--start-year", type=int, required=True)
    parser.add_argument("--end-year", type=int, required=True)
    parser.add_argument("--output", type=Path, default=Path("statistik_bps.xlsx"))
    args = parser.parse_args()
    category = next(item["name"] for item in SUBJECT_CATEGORIES if item["id"] == args.category_id)
    subjects = fetch_subject_types(args.api_key, args.category_id)
    subject = next((item["name"] for item in subjects if item["id"] == args.subject_id), "")
    if not subject:
        parser.error("subject-id tidak ditemukan pada kategori yang dipilih.")
    tables = fetch_tables(args.api_key, args.subject_id)
    table = next((item for item in tables if item["id"] == args.table_id), None)
    if not table:
        parser.error("table-id tidak ditemukan pada jenis data yang dipilih.")
    rows, missing = scrape_table_by_year(
        args.api_key, category, subject, args.table_id, args.title or table["title"], table["source"],
        args.start_year, args.end_year,
        progress=lambda message, _: print(message),
    )
    save_xlsx(rows, args.output, {"Subjek": category, "Jenis data": subject, "Nama tabel": table["title"],
                                  "Periode diminta": f"{args.start_year}–{args.end_year}", "Tahun tanpa data": missing})
    print(f"Baris data: {len(rows)}; tahun tanpa data: {missing}; file: {args.output.resolve()}")


if __name__ == "__main__":
    main()

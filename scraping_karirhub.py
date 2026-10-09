"""Ambil lowongan publik Karirhub berdasarkan filter provinsi dan simpan ke XLSX."""

from __future__ import annotations

import argparse
import re
import sys
import time
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable
from urllib.parse import parse_qs, quote, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


KARIRHUB_HOST = "karirhub.kemnaker.go.id"
KARIRHUB_WEB = f"https://{KARIRHUB_HOST}"
API_ROOT = "https://api.kemnaker.go.id/karirhub"
API_BASE = f"{API_ROOT}/vacancy/v2"
LISTING_API = f"{API_BASE}/published-industrial-vacancies"
PROVINCES_API = f"{API_ROOT}/onboarding/v2/provinces"
EDUCATIONS_API = f"{API_ROOT}/v1/educations"
JOB_TYPES_API = f"{API_ROOT}/v1/job-types"
DEFAULT_URL = (
    f"{KARIRHUB_WEB}/lowongan-dalam-negeri/lowongan"
    "?locations%5B0%5D%5Bid%5D=75ca55c9-7b5a-43fb-837d-327b3e918d09"
    "&locations%5B0%5D%5Btype%5D=province"
    "&locations%5B0%5D%5Blabel%5D=Sumatera%20Selatan"
)
PAGE_SIZE = 18
DEFAULT_DELAY = 0.5
SOURCE_CATEGORIES = [
    {"value": "all", "label": "Semua sumber"},
    {"value": "top", "label": "Top (Toploker)"},
    {"value": "glints", "label": "Glints"},
    {"value": "jobstreet", "label": "JobStreet"},
    {"value": "other", "label": "Lainnya"},
]
COLUMNS = [
    "Jabatan",
    "Perusahaan",
    "Deskripsi Pekerjaan",
    "Kualifikasi",
    "Gaji",
    "Lokasi",
    "Tanggal Posting",
    "Sumber Lowongan",
    "URL Lowongan",
]
ProgressCallback = Callable[[str, int], None]


def get_json(
    session: requests.Session,
    url: str,
    params: dict[str, object] | None = None,
    timeout: int = 30,
) -> dict:
    response = session.get(url, params=params, timeout=timeout)
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, dict) or "data" not in payload:
        raise ValueError(f"Format respons Karirhub tidak sesuai: {url}")
    return payload


def fetch_filter_options(timeout: int = 15) -> dict[str, list[dict[str, str]]]:
    """Ambil dropdown provinsi, pendidikan, dan tipe kerja dari API publik Karirhub."""
    with requests.Session() as session:
        session.headers.update({"Accept": "application/json", "User-Agent": "LowonganDataUI/1.0"})
        provinces = get_json(session, PROVINCES_API, timeout=timeout).get("data") or []
        educations = get_json(session, EDUCATIONS_API, timeout=timeout).get("data") or []
        job_types = get_json(session, JOB_TYPES_API, timeout=timeout).get("data") or []

    return {
        "provinces": sorted(
            [{"id": str(item["id"]), "name": str(item["name"])} for item in provinces if item.get("id") and item.get("name")],
            key=lambda item: item["name"].casefold(),
        ),
        "educations": sorted(
            [{"id": str(item["id"]), "name": str(item["name"]), "level": item.get("level", 999)} for item in educations if item.get("id") and item.get("name")],
            key=lambda item: int(item["level"] or 999),
        ),
        "job_types": sorted(
            [{"id": str(item["id"]), "name": str(item["name"])} for item in job_types if item.get("id") and item.get("name")],
            key=lambda item: item["name"].casefold(),
        ),
        "sources": SOURCE_CATEGORIES.copy(),
    }


def location_filter_from_url(url: str) -> tuple[str, str, str]:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != KARIRHUB_HOST:
        raise ValueError(f"URL harus berasal dari https://{KARIRHUB_HOST}.")
    if parsed.path.rstrip("/") != "/lowongan-dalam-negeri/lowongan":
        raise ValueError("URL harus menunjuk ke halaman daftar lowongan dalam negeri Karirhub.")

    query = parse_qs(parsed.query)
    location_id = query.get("locations[0][id]", [""])[0]
    location_type = query.get("locations[0][type]", [""])[0]
    location_label = query.get("locations[0][label]", [""])[0]
    if not location_id or not location_type or not location_label:
        raise ValueError("URL perlu memuat filter lokasi id, type, dan label.")
    if location_type != "province":
        raise ValueError("Filter lokasi pada URL harus bertipe province.")
    return location_id, location_type, location_label


def get_listing_rows(
    session: requests.Session,
    location: tuple[str, str, str],
    page_size: int = PAGE_SIZE,
    delay: float = DEFAULT_DELAY,
    progress: ProgressCallback | None = None,
) -> list[dict]:
    location_id, location_type, location_label = location
    rows: list[dict] = []
    seen_ids: set[str] = set()
    page = 1

    while True:
        params: dict[str, object] = {
            "page": page,
            "limit": page_size,
            "locations[0][id]": location_id,
            "locations[0][type]": location_type,
            "locations[0][label]": location_label,
        }
        payload = get_json(session, LISTING_API, params=params)
        page_rows = payload.get("data") or []
        if not isinstance(page_rows, list) or not page_rows:
            break

        for row in page_rows:
            if not isinstance(row, dict):
                continue
            row_id = str(row.get("id") or "")
            if row_id and row_id not in seen_ids:
                seen_ids.add(row_id)
                rows.append(row)

        meta = payload.get("meta") or {}
        last_page = meta.get("last_page", meta.get("lastPage", page))
        try:
            last_page = int(last_page)
        except (TypeError, ValueError):
            last_page = page
        if progress:
            progress(f"Membaca lowongan {location_label} — halaman {page}/{last_page}", min(20, 5 + int(15 * page / max(last_page, 1))))
        if page >= last_page:
            break

        page += 1
        if delay:
            time.sleep(delay)

    return rows


def source_matches(platform_name: object, source: str) -> bool:
    if source in {"", "all"}:
        return True
    name = str(platform_name or "").casefold()
    known = {"top": "toploker", "glints": "glints", "jobstreet": "jobstreet"}
    if source == "other":
        return not any(value in name for value in known.values())
    expected = known.get(source)
    return bool(expected and expected in name)


def filter_listings(
    listings: list[dict],
    job_type_id: str = "",
    education_id: str = "",
    source: str = "all",
    max_jobs: int = 0,
) -> list[dict]:
    filtered = []
    for item in listings:
        if job_type_id and str(item.get("job_type_id") or "") != job_type_id:
            continue
        if education_id and str(item.get("education_id") or "") != education_id:
            continue
        if not source_matches(item.get("platform_name"), source):
            continue
        filtered.append(item)
        if max_jobs and len(filtered) >= max_jobs:
            break
    return filtered


def clean_text(value: object) -> str:
    if value is None:
        return ""
    soup = BeautifulSoup(str(value), "html.parser")
    text = soup.get_text("\n", strip=True)
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def format_amount(value: object) -> str:
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return ""
    if amount == amount.to_integral_value():
        return f"{int(amount):,}".replace(",", ".")
    return f"{amount:,.2f}".rstrip("0").rstrip(".").replace(",", ".")


def format_salary(detail: dict) -> str:
    if detail.get("show_salary") is False:
        return "Dirahasiakan"

    minimum = detail.get("min_salary_amount")
    maximum = detail.get("max_salary_amount")
    if minimum is None:
        minimum = (detail.get("min_salary") or {}).get("amount")
    if maximum is None:
        maximum = (detail.get("max_salary") or {}).get("amount")
    if minimum in (None, "") and maximum in (None, ""):
        return "Dirahasiakan"
    try:
        if Decimal(str(minimum or 0)) == 0 and Decimal(str(maximum or 0)) == 0:
            return "Dirahasiakan"
    except (InvalidOperation, TypeError, ValueError):
        return "Dirahasiakan"

    currency = str(detail.get("min_salary_currency") or detail.get("max_salary_currency") or "IDR")
    currency_label = "Rp" if currency.upper() == "IDR" else currency.upper()
    minimum_text = format_amount(minimum)
    maximum_text = format_amount(maximum)
    if minimum_text and maximum_text and minimum_text == maximum_text:
        return f"{currency_label} {minimum_text}"
    if minimum_text and maximum_text:
        return f"{currency_label} {minimum_text} – {currency_label} {maximum_text}"
    amount = minimum_text or maximum_text
    return f"{currency_label} {amount}" if amount else "Dirahasiakan"


def qualification_text(detail: dict) -> str:
    parts: list[str] = []
    special = clean_text(detail.get("qualification"))
    if special:
        parts.append(special)

    education = detail.get("min_education") or {}
    education_name = education.get("name") if isinstance(education, dict) else ""
    if education_name:
        parts.append(f"Pendidikan minimum: {education_name}")

    years = detail.get("min_year_experience")
    if years not in (None, ""):
        parts.append(f"Pengalaman minimum: {years} tahun")

    age_min = detail.get("min_age")
    age_max = detail.get("max_age")
    if age_min is not None and age_max is not None:
        parts.append(f"Rentang usia: {age_min}–{age_max} tahun")
    elif age_min is not None:
        parts.append(f"Usia minimum: {age_min} tahun")
    elif age_max is not None:
        parts.append(f"Usia maksimum: {age_max} tahun")

    gender = detail.get("gender")
    gender_labels = {"male": "Laki-laki", "female": "Perempuan", "both": "Laki-laki / Perempuan"}
    if gender:
        parts.append(f"Jenis kelamin: {gender_labels.get(str(gender).lower(), gender)}")

    skills = detail.get("skills") or []
    if isinstance(skills, list) and skills:
        skills_text = ", ".join(clean_text(skill) for skill in skills if skill)
        if skills_text:
            parts.append(f"Keterampilan: {skills_text}")
    return "\n".join(parts)


def make_job_url(title: str, job_id: str) -> str:
    normalized = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", normalized.lower()).strip("-")
    return f"{KARIRHUB_WEB}/lowongan-dalam-negeri/lowongan/{quote(slug, safe='-')}-{job_id}"


def parse_posted_date(value: object) -> datetime | str:
    if not value:
        return ""
    text = str(value)
    for date_format in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            pass
    return text


def detail_to_row(listing: dict, detail: dict, province_label: str) -> list[object]:
    employer = detail.get("employer") or {}
    region = detail.get("region") or {}
    platform = detail.get("platform") or {}
    if not isinstance(employer, dict):
        employer = {}
    if not isinstance(platform, dict):
        platform = {}
    title = str(detail.get("title") or listing.get("title") or "").strip()
    company = employer.get("name") or listing.get("company_name") or ""
    row_id = str(detail.get("id") or listing.get("id") or "")
    location_name = region.get("name") if isinstance(region, dict) else ""

    return [
        title,
        company,
        clean_text(detail.get("description")),
        qualification_text(detail),
        format_salary(detail),
        location_name or province_label,
        parse_posted_date(detail.get("created_at") or listing.get("created_at")),
        str(platform.get("name") or listing.get("platform_name") or ""),
        make_job_url(title, row_id) if row_id else "",
    ]


def get_job_details(session: requests.Session, job_id: str) -> dict:
    payload = get_json(session, f"{API_BASE}/published-industrial-vacancies/{job_id}")
    detail = payload.get("data")
    if not isinstance(detail, dict):
        raise ValueError(f"Detail lowongan {job_id} tidak ditemukan di respons Karirhub.")
    return detail


def scrape_jobs_filtered(
    province_id: str,
    province_label: str,
    job_type_id: str = "",
    education_id: str = "",
    source: str = "all",
    page_size: int = PAGE_SIZE,
    max_jobs: int = 0,
    delay: float = DEFAULT_DELAY,
    progress: ProgressCallback | None = None,
) -> tuple[list[list[object]], int]:
    if not province_id or not province_label:
        raise ValueError("Pilih provinsi sebelum mengambil lowongan.")
    if source not in {"all", "top", "glints", "jobstreet", "other"}:
        raise ValueError("Pilihan sumber lowongan tidak valid.")
    if not 1 <= page_size <= 100 or max_jobs < 0 or delay < 0:
        raise ValueError("Ukuran halaman, batas jumlah, atau jeda tidak valid.")

    location = (province_id, "province", province_label)
    session = requests.Session()
    session.headers.update({"Accept": "application/json", "User-Agent": "LowonganDataUI/1.0"})
    try:
        listings = get_listing_rows(session, location, page_size, delay, progress)
        selected = filter_listings(listings, job_type_id, education_id, source, max_jobs)
        results: list[list[object]] = []
        detail_failures = 0
        total = len(selected)

        for index, listing in enumerate(selected, start=1):
            job_id = str(listing.get("id") or "")
            if not job_id:
                continue
            if index > 1 and delay:
                time.sleep(delay)
            try:
                detail = get_job_details(session, job_id)
            except (requests.RequestException, ValueError) as error:
                detail_failures += 1
                print(f"Detail tidak terbaca untuk {job_id}: {error}", file=sys.stderr)
                detail = {}
            results.append(detail_to_row(listing, detail, province_label))
            if progress:
                percent = 20 + int(index / max(total, 1) * 78)
                progress(f"Memuat detail lowongan {index}/{total}: {listing.get('title', job_id)}", percent)

        return results, detail_failures
    finally:
        session.close()


def scrape_jobs(
    url: str,
    page_size: int = PAGE_SIZE,
    max_jobs: int = 0,
    delay: float = DEFAULT_DELAY,
    progress: ProgressCallback | None = None,
) -> tuple[list[list[object]], int]:
    location_id, _, location_label = location_filter_from_url(url)
    return scrape_jobs_filtered(
        province_id=location_id,
        province_label=location_label,
        page_size=page_size,
        max_jobs=max_jobs,
        delay=delay,
        progress=progress,
    )


def save_xlsx(rows: list[list[object]], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lowongan"
    sheet.sheet_view.showGridLines = False
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill(fill_type="solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30
    sheet.auto_filter.ref = f"A1:I{max(len(rows) + 1, 1)}"
    widths = [28, 30, 60, 55, 28, 35, 22, 22, 55]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[6].number_format = "yyyy-mm-dd hh:mm"
        if row[8].value:
            row[8].hyperlink = row[8].value
            row[8].style = "Hyperlink"

    if rows:
        table = Table(displayName="TabelLowonganKarirhub", ref=f"A1:I{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Ambil lowongan Karirhub berdasarkan URL provinsi dan ekspor ke XLSX.")
    parser.add_argument("--url", default=DEFAULT_URL, help="URL daftar lowongan Karirhub dengan filter provinsi.")
    parser.add_argument("--output", type=Path, default=Path("lowongan_karirhub.xlsx"))
    parser.add_argument("--page-size", type=int, default=PAGE_SIZE)
    parser.add_argument("--max-jobs", type=int, default=0, help="0 berarti semua lowongan yang cocok.")
    parser.add_argument("--delay", type=float, default=DEFAULT_DELAY)
    args = parser.parse_args()
    if not 1 <= args.page_size <= 100 or args.max_jobs < 0 or args.delay < 0:
        parser.error("page-size harus 1–100; max-jobs dan delay tidak boleh negatif.")

    try:
        rows, failures = scrape_jobs(args.url, args.page_size, args.max_jobs, args.delay, lambda message, _: print(message))
        save_xlsx(rows, args.output)
    except (requests.RequestException, ValueError, OSError) as error:
        parser.exit(1, f"Gagal mengambil atau menyimpan data Karirhub: {error}\n")

    print(f"Jumlah lowongan: {len(rows)}")
    if failures:
        print(f"Detail gagal diambil: {failures} (kolom detail terkait dibiarkan kosong)")
    print(f"File XLSX tersimpan di: {args.output.resolve()}")


if __name__ == "__main__":
    main()
